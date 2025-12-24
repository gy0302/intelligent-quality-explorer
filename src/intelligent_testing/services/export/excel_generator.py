"""
MVP核心：Excel生成器
将测试用例导出为Excel格式
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any
from pathlib import Path

from intelligent_testing.models.test_case import TestCase


class ExcelGenerator:
    """Excel生成器（MVP版本）"""

    def __init__(self):
        self.styles = self._define_styles()

    def _define_styles(self) -> Dict[str, Any]:
        """定义单元格样式（简化版）"""
        return {
            "header": {
                "font": Font(bold=True, color="FFFFFF", size=12),
                "fill": PatternFill(
                    start_color="4F81BD",  # 蓝色
                    end_color="4